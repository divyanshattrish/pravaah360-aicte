import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

print("Testing flowchart shapes...")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Test shapes: Oval, Diamond, Rounded Rectangle, Down Arrow, Can (database cylinder)
s1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1), Inches(1), Inches(1.5), Inches(0.5))
s2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(2.0), Inches(0.8))
s3 = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(1), Inches(3.2), Inches(1.8), Inches(1.0))
s4 = slide.shapes.add_shape(MSO_SHAPE.CAN, Inches(1), Inches(4.5), Inches(1.5), Inches(1.0))
s5 = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(1.5), Inches(1.6), Inches(0.3), Inches(0.3))

print("All shapes supported successfully!")
